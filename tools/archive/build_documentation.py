from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Subtitle',11),('Heading 1',14),('Heading 2',11)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0)
    st.paragraph_format.space_after=Pt(7)

def p(text,style=None):
    return doc.add_paragraph(text,style)

def h(text):
    doc.add_heading(text,1)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,values):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE8ED' if ri==0 else 'F4F7F8' if ri%2==0 else 'FFFFFF');props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
    p('') .paragraph_format.space_after=Pt(0)

p('Trial Limit System Development Summary','Title')
p('Work completed on 15 October 2026','Subtitle')
p('I built a local Python tool and browser dashboard to check whether enrolled students plus manually simulated trial students fit within a capacity of 100 for a selected age batch and session. The current version is a working capacity-checking prototype for testing; it does not create trial bookings.')

h('Data source and initial understanding')
p('Sarthak from TechOps provided the enrollment data and gave me an initial overview of the tables, important fields and distinctions in the data. I used that explanation to understand the supplied enrollment export and develop the counting logic. The application currently reads local files; it has no live database connection.')
p('The enrollment source is academy___enrollment_data___formatted_report_2026-10-15T15_16_51.527944248+05_30.csv. The supplied snapshot contains 14,741 enrollment records. A record is not necessarily a unique student, and the file includes historical enrollments.')
p('Greenfield Batch & Coach details(reference for the batch timings).xlsx was supplied as a separate reference for centers, batch days, timings and divisions. The application reads its MAIN sheet. Its labels remain reference labels and do not replace center names in the enrollment CSV.')
p('The later reference file is an Excel workbook, rather than a second enrollment CSV. It helped me understand how the batches operate: which weekdays they run, their time intervals, and how age or proficiency divisions can share a session. The CSV answers who is enrolled; the workbook helps explain which scheduled batch should be checked. Together they support a session-specific limit instead of treating every student at a center as one batch.')
p('Both files remain strictly read-only. Coaches and data-center teams may maintain the source data. The tool preserves their identifiers, capitalization, spacing, missing values and inconsistencies. Dates, booleans and timings are interpreted only in memory; reports are written separately.')

h('Work completed today')
p('The first implementation was a command-line application with configurable capacity, a selectable reference date, unique-student counts, manual trial simulation and separate summary and data-quality exports. Its original capacity groups used Academy ID, location, sport, age category and course.')
p('A local frontend was then added to make testing easier. It shows exact CSV sport, academy, center and age selections, enrollment and trial counts, occupancy, remaining places and capacity decisions. The latest update adds workbook batch references and separate limits for age batches at specific weekdays and times.')
p('Where a workbook center label differs from the CSV, the operator selects the corresponding reference row and confirms the temporary link. No name is changed, no alias is guessed and the link is not written into either source file.')

h('Important data distinctions')
p('Player DB ID identifies a student and is preserved as a string. Enrollment ID identifies an enrollment record. Academy ID is retained as the grouping identifier associated with Academy Name; Sports remains a separate source field. Slots provides the scheduled weekday and time. Workbook batch descriptions provide reference divisions such as U6, U8 and U10.')

doc.add_page_break()
h('How the current batch check works')
p('The operator selects a session date and a CSV academy, center, sport and age category, then selects and confirms the corresponding workbook timing reference. The program counts unique Player DB IDs whose Active value is True, whose Start Date and End Date include the session date, and whose Slots match that weekday and exact time interval.')
p('Each separate age batch has a default limit of 100. U8 and U10 can each have 100 at the same center and time. For a workbook division that specifies an age without a separate course division, unique students across courses share that age batch limit. A course-only subset cannot receive approval for the whole age batch.')
p('The limit is configurable in config.json. Workbook Total Capacity and Actual Capacity values are not used as limits or enrollment counts in this version. Trial counts are entered manually because no live trial-booking source has been connected.')

h('Example of the capacity rule')
p('The figures below are fictional test examples and are not actual student counts.')
table(['Age batch and session','Enrolled','Trials','Total','Result'],[
    ['Saturday 6–7 PM U6','70','10','80','Available\n20 places left'],
    ['Saturday 6–7 PM U8','95','5','100','Full\nNo additional places'],
    ['Saturday 7–8 PM U10','98','5','103','Over Capacity\n3 excess students'],
],[2.05,.65,.6,.6,2.85-0.0])
p('Total occupancy = unique enrolled students + current simulated trials. Remaining places = max(0, capacity − total occupancy). Excess students = max(0, total occupancy − capacity). Proposed additional trials are allowed only when data is complete and total occupancy + proposed trials is at most 100.')
p('A proposal check does not add students to saved occupancy or create a booking. If current occupancy is already above 100, the tool reports the excess and rejects additional trials. Without an explicit simulated trial count, combined occupancy and trial availability remain unconfirmed.')

h('Validation and source protection')
p('All 26 automated tests passed, covering duplicate students, separate age batches, dates, invalid inputs, missing data, unconfirmed links and capacity thresholds. Browser checks verified the workflow and that incomplete results withhold approval.')
p('SHA-256 comparisons verified that both original source files remained unchanged during testing. Generated summaries and data-quality reports are separate downloads. The earlier command-line application also writes timestamped exports to the outputs folder.')

h('Current limitations and running the tool')
p('Missing IDs, invalid Slots or dates, ambiguous batch divisions and unresolved links make affected results incomplete. Counts are then provisional and approval is withheld. There is no live attendance or trial-booking integration. Reference links and simulations are temporary, and checks use exact intervals rather than combining overlapping sessions.')
p('In VS Code PowerShell, run python frontend.py from the project folder and open http://127.0.0.1:8765/. Keep the terminal open. Run python -m unittest -v for tests; see README.md for detailed instructions.')

doc.core_properties.title='Trial Limit System Development Summary'
doc.core_properties.subject='Data provenance and work completed on 15 October 2026'
doc.core_properties.author=''
output=OUT/'Trial_Limit_System_Development_Summary_2026-10-15.docx'
doc.save(output)
print(output)
